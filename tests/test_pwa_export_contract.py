"""Verifica el round-trip completo del contrato del ZIP: el paquete que arma la
lógica de exportación de la PWA (pwa/js/export-zip.js, corrida tal cual con
Node — no una reimplementación en Python que se pueda desalinear) pasa la
validación real de `receipt_collections.import_collection_from_zip` y se
importa correctamente. Ese es el criterio de aceptación acordado para la PWA:
exportar desde el celular -> importar en el PC -> la colección aparece con sus
boletas.

También cubre el caso motivador de que una boleta se nombre con una descripción
con espacios (ej. "almuerzo con cliente dos personas", vía ParserBoletasDB.
sanitizeReceiptName) — ese nombre de archivo tiene que sobrevivir intacto hasta
la columna Comments del Excel de rendición.

Se salta si no hay Node disponible o si no se corrió `npm install` en pwa/
(nada del resto de la suite de la app de PC depende de Node)."""

import json
import shutil
import subprocess
from pathlib import Path
from unittest.mock import patch

import pytest
import yaml

import excel_writer
import main
import receipt_collections as rc
from extract import ExtractionResult
from validate import validate_extraction

PWA_DIR = Path(__file__).resolve().parent.parent / "pwa"
BUILD_SCRIPT = PWA_DIR / "scripts" / "build_test_zip.js"
PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = PROJECT_ROOT / "plantilla" / "Form - Chile Expense Report - Peru Travel Junio 2026.xlsx"

DESCRIBED_RECEIPT_NAME = "almuerzo con cliente dos personas"  # ya saneado (espacios colapsados)


@pytest.fixture(autouse=True)
def _isolated_collections_root(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "_collections_root", lambda: tmp_path / "colecciones")


def _node_available() -> bool:
    return shutil.which("node") is not None and (PWA_DIR / "node_modules" / "jszip").exists()


def _build_pwa_zip(output_zip: Path, collection_name: str = "Coleccion Prueba PWA") -> None:
    result = subprocess.run(
        ["node", str(BUILD_SCRIPT), str(output_zip), collection_name],
        cwd=PWA_DIR,
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, f"build_test_zip.js falló:\n{result.stdout}\n{result.stderr}"
    assert output_zip.exists()


@pytest.mark.skipif(not _node_available(), reason="node y/o pwa/node_modules/jszip no disponibles")
def test_pwa_exported_zip_passes_real_importer(tmp_path):
    output_zip = tmp_path / "pwa-export.zip"
    _build_pwa_zip(output_zip)

    summary = rc.import_collection_from_zip(output_zip)

    assert summary.name == "Coleccion Prueba PWA"
    assert summary.n_receipts == 3
    assert sorted(p.name for p in rc.list_receipts(summary.slug)) == [
        f"{DESCRIBED_RECEIPT_NAME}.jpg",
        "boleta_clp.jpg",
        "boleta_usd.jpg",
    ]

    # El coleccion.json reconstruido queda en el esquema local estándar, igual
    # que cualquier otra colección (ver receipt_collections.py).
    metadata = json.loads(rc._metadata_path(summary.slug).read_text(encoding="utf-8"))
    assert set(metadata.keys()) == {"version", "name", "created_at", "last_generated_at"}
    assert metadata["version"] == rc.CURRENT_METADATA_VERSION


@pytest.mark.skipif(not _node_available(), reason="node y/o pwa/node_modules/jszip no disponibles")
def test_pwa_receipt_name_with_spaces_reaches_comments_column(tmp_path):
    """Round-trip completo: nombre con espacios puesto en la PWA -> importado
    en el PC -> generado el Excel -> aparece tal cual en la columna Comments."""
    output_zip = tmp_path / "pwa-export.zip"
    _build_pwa_zip(output_zip)
    summary = rc.import_collection_from_zip(output_zip)
    files = rc.list_receipts(summary.slug)

    config = yaml.safe_load(TEMPLATE_PATH.parent.parent.joinpath("config.yaml").read_text(encoding="utf-8"))
    config["excel"]["template_path"] = str(TEMPLATE_PATH)

    def fake_process_file(file_path, cfg):
        result = ExtractionResult(
            date="2026-06-10",
            currency="CLP",
            amount=18500.0,
            expense_type="Travel - Meals",
            confidence={"overall": 0.95, "amount": 0.95, "currency": 0.95, "date": 0.95, "expense_type": 0.95},
        )
        return result, validate_extraction(result, cfg)

    with patch.object(main, "process_file", fake_process_file):
        summary_result = main.process_all(files, config)

    output_path = tmp_path / "rendicion.xlsx"
    excel_writer.write_expense_report(summary_result.rows_to_write, config, output_path)

    import openpyxl

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Expense Report"]
    comments = [ws.cell(row=r, column=9).value for r in range(9, 12) if ws.cell(row=r, column=9).value]
    assert DESCRIBED_RECEIPT_NAME in comments
